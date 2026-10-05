from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from pathlib import Path

ROOT=Path(__file__).parent
doc=Document()
sec=doc.sections[0]
sec.top_margin=sec.bottom_margin=Inches(.65)
sec.left_margin=sec.right_margin=Inches(.65)
sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
for name in ['Normal','Title','Heading 1','Heading 2']:
 s=doc.styles[name]; s.font.name='Calibri'; s.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(10)
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Title'].font.size=Pt(25)
doc.styles['Heading 1'].font.size=Pt(17)
doc.styles['Heading 2'].font.size=Pt(12)

def p(t): doc.add_paragraph(t)
def h(t): doc.add_heading(t,1)
def table(headers, rows, widths):
 t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
 for c,w in zip(t.columns,widths): c.width=Inches(w)
 for c,v in zip(t.rows[0].cells,headers): c.text=v
 for row in rows:
  for c,v in zip(t.add_row().cells,row): c.text=v
 for i,row in enumerate(t.rows):
  for j,c in enumerate(row.cells):
   c.width=Inches(widths[j]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=c._tc.get_or_add_tcPr()
   bd=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); bd.append(e)
   pr.append(bd)
   mar=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge); e.set(qn('w:w'),'90'); e.set(qn('w:type'),'dxa'); mar.append(e)
   pr.append(mar)
   shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E7EEF5' if i==0 else ('F5F5F5' if i%2==0 else 'FFFFFF')); pr.append(shade)
   for para in c.paragraphs:
    para.paragraph_format.space_after=Pt(2)
    para.paragraph_format.space_before=Pt(2)
    for r in para.runs: r.font.size=Pt(9); r.bold=i==0
  rr=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(rr)
 repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
 doc.add_paragraph().paragraph_format.space_after=Pt(0)
 return t

doc.add_paragraph('YouTube Audio Story Website Project Tracker','Title')
p('Updated 5 October 2026 | Version 4 | Project stage: local prototype; multi-user roadmap added')
p('We are building a mobile-friendly website that recommends YouTube audio stories based on a listener’s taste and current need. This tracker breaks the work into small steps, with progress shown beside every step. The website is for many listeners, not personal use. A local prototype now exists. Version 3 revised the recommendation logic and re-audited milestone C only; Version 4 adds milestone F, the multi-user roadmap; milestones A, B, D and E still show their Version 2 status and need a separate review.')
h('Agreed direction')
p('Start with a website. Collect favorite story links and explicit feedback, recommend a few suitable stories, and provide a route to YouTube playback. Treat the recommendation process as small problems with clear inputs and outputs.')
p('Proposed first-version architecture: Cloudflare Pages hosts the website and curated catalog. The browser filters and ranks stories and saves preferences and feedback locally. AI enriches the catalog during preparation; live AI calls are deferred. Hosting setup and AI processing costs must be checked before implementation.')
h('Overall progress')
table(['Milestone','Completed steps','Progress'],[
 ('A Define listener needs','0 / 4','0% | Not started'),('B Prepare and enrich the catalog','0 / 7','0% | Not started'),('C Define recommendation rules','2 / 4','50% | In progress'),('D Build the first website','0 / 6','0% | Not started'),('E Validate and release','0 / 4','0% | Not started'),('F Multi-user platform','0 / 8','0% | Not started')],[3.35,1.25,2.37])
p('Implementation checklist: 2 of 33 steps marked complete (C2, C4). Other milestones have not been re-audited since Version 2, so the true total is likely higher. Percent complete = completed steps / total steps × 100. This counts deliverables, not effort or time.')
h('How to maintain the tracker')
p('Replace each status with Not started, In progress, Blocked, or Done. Mark Done only when the completion evidence exists. Update milestone counts and percentages after each work session. Assign owners and dates before starting; neither is committed yet.')
p('First action: choose the initial language and listener group, then record three examples of the listening problem.')

doc.add_page_break()
h('A Define listener needs')
p('Small problem: what minimum information lets us make a useful first recommendation? Output: an agreed listener profile and first-version scope.')
table(['ID','Small step and completion evidence','Progress'],[
 ('A1','Choose one initial language and listener group. Done when both are recorded in the decision log.','Not started'),
 ('A2','Collect three listening situations and the difficulty in each. Done when concrete examples are written.','Not started'),
 ('A3','Choose the starting questions: favorite links, genre, narrator preferences, mood and duration. Done when the minimum fields are agreed.','Not started'),
 ('A4','Approve the first-version scope and exclusions. Done when the core user journey and success criteria are recorded.','Not started')],[.45,5.15,1.37])
h('B Prepare the story catalog')
p('Small problem: what do we need to know about each story to compare it fairly? Output: a small, usable catalog with consistent attributes.')
table(['ID','Small step and completion evidence','Progress'],[
 ('B1','Define story fields: video link, title, language, genre, narrator, duration, mood, series order and content labels. Done when a sample record is complete.','Not started'),
 ('B2','Investigate permitted YouTube discovery and playback options, API access and quotas. Done when the chosen approach and constraints are documented.','Not started'),
 ('B3','Curate an initial catalog. Proposed target: 30 stories in the chosen language. Done when links and essential fields are checked.','Not started'),
 ('B4','Define missing-data and unavailable-video handling. Done when incomplete records and broken links have a clear treatment.','Not started')],[.45,5.15,1.37])
p('Dependency: A1 before catalog selection; B1 and B2 before B3. Catalog size is a proposed starting target, not an agreed commitment.')

doc.add_page_break()
h('AI role in the first version')
p('AI works during catalog preparation. It turns available descriptions or permitted transcripts into structured tags for themes, setting, mood and storytelling style. A person reviews those tags before they enter the website catalog. The browser then uses ordinary scoring code to recommend verified stories.')
p('Flow: story information → AI tagging → human review → enriched catalog → browser filters and ranking → three recommendations → YouTube playback. Explicit listener feedback updates the local taste profile and later rankings.')
p('Version 3 adds two more offline AI uses, both prepared before the website loads: text embeddings that measure plot and mood similarity between stories, and an optional LLM score with a one-sentence reason for the top candidates. The website itself calls no AI service. See Revised recommendation logic.')
table(['ID','Small step and completion evidence','Progress'],[
 ('B5','Define the AI tagging schema and input sources. Done when a sample includes themes, setting, mood, source evidence and unknown fields, and an AI tool and cost approach are chosen.','Not started'),
 ('B6','Tag the initial catalog with AI and review every record. Done when supported tags are approved, unsupported claims removed, and original links retained.','Not started'),
 ('B7','Export the reviewed catalog as a website data file. Done when each video ID is unique and the browser can load validated tags without live AI calls.','Not started')],[.45,5.15,1.37])
h('Evidence and controls')
p('AI must use the supplied story information and preserve verified video IDs. Unknown attributes stay unknown. Voice quality, narration pace and background music require audio evidence or human review; text descriptions alone do not establish them. Recommendation explanations use saved matching attributes rather than invented facts or links.')
h('Future live AI option')
p('A later version could interpret requests such as “something like yesterday’s story, but less frightening.” Proposed flow: browser → backend → AI service → validated catalog candidates → strict filters → recommendations. This is deferred work and is excluded from the 25-step first-version checklist.')
p('Before adding live AI, decide the provider, budget, backend, consent and data handling. Keep API credentials on the backend. Free website hosting does not make AI processing free. The first version needs no per-visit AI request or paid database; local data does not sync across devices.')
p('Dependencies: B1, B2 and B3 before B5; B5 before B6; B4 and B6 before B7; B7 before C3 and D4. The first AI experiment uses the proposed 30-story catalog and three example listener profiles.')

doc.add_page_break()
h('C Define recommendation rules')
p('Small problem: which eligible stories best fit this listener today? Output: explainable ranking rules that can be checked before building the interface.')
table(['ID','Small step and completion evidence','Progress'],[
 ('C1','Define hard filters such as language, maximum duration, hidden channels and already-heard stories. Done when each filter has an example.','In progress: genre, verified-duration and hidden-story filters work; heard stories sort last; channel hiding not built'),
 ('C2','Define browser scoring over reviewed AI tags using taste, current need and discovery value. Done when weights, missing-data rules and attribute-based explanations are recorded.','Done: two-stage model below; recommendations.js, README and unit tests'),
 ('C3','Try three listener profiles against the enriched catalog. Done when each receives three verified, eligible suggestions with reasons and differences from genre-only matching are reviewed.','In progress: one real profile tested with leave-one-out; two more needed'),
 ('C4','Define feedback updates and a fallback for sparse matches. Done when likes, dislikes and already-heard feedback produce expected changes.','Done: like, not-for-me, heard and restore covered by tests; sparse matches fall back to exploration')],[.45,5.15,1.37])
p('Version 2 proposed score = weighted taste match + weighted current-need match + weighted discovery value. Version 3 replaces it with the probabilistic model below.')
doc.add_heading('Revised recommendation logic (Version 3)',2)
p('Each story receives a probability that this listener will enjoy it, P(enjoy). Hard filters apply first: genre, verified duration (stories without a verified duration are excluded from duration filters) and hidden stories. Already-heard stories sort after unheard ones.')
table(['Term','What it measures','How it is calculated'],[
 ('Stage 1','Probability score','logit P(enjoy i) = content(i) + behavior(i) + explicit(i) + γ · emb(i)'),
 ('content','Match on series, author, genre and channel','Each attribute holds a Beta belief starting at the listener’s overall rate. Evidence for: confirmed favorite 3, Like 2. Against: Not for me 2. Combined on the log-odds scale with weights series 1.0, author 0.8, genre 0.5, channel 0.3, scaled by inverse document frequency so common genres such as suspense count less.'),
 ('behavior','Weak interest from YouTube history','Return days and recency, fitted offline with Bayesian logistic regression and capped at ±0.25 log-odds. Gap-based completion and skip estimates are diagnostics only and do not affect ranking.'),
 ('explicit','Listener feedback','Like adds 2 log-odds. Floors and ceilings: confirmed favorite at least 95%, liked at least 85%, Not for me at most 5%.'),
 ('emb','Plot and mood similarity (content-based filtering with text embeddings)','Weighted average over the 10 most similar stories: Σ w·(r − r̄) / Σ w, where w = max(0, cosine − baseline). r is 1 for loved or liked stories, −1 for Not for me, and at most ±0.25 for stories only opened. Embeddings come from multilingual-E5 on title and description.'),
 ('Stage 2','LLM rerank of the top 20 unheard candidates','logit P_final = logit P(enjoy) + δ · logit(s_LLM). s_LLM is an LLM fit score prepared offline with a one-sentence reason. Cards with s_LLM ≥ 0.6 show an “AI-reviewed” reason.'),
 ('Discovery','New stories beyond history','Thompson sampling from the Beta beliefs, seeded by date so the list is stable within a day, plus a series-diversity pass.')],[1.0,1.9,4.07])
p('Weights: γ has prior N(1, 0.5) and δ has prior N(0.5, 0.5). Both are fitted by fit_preferences.py and kept only if leave-one-out ranking of loved stories does not get worse; otherwise they are set to 0. δ is fitted only from stories the listener labelled, and only once at least five labelled stories have LLM scores, because the LLM was shown the favorites.')
p('Explanations name the strongest contributor: listener feedback, an AI-reviewed reason, repeated openings, plot and mood similarity to a loved story, or a series, author or genre match. Stories below 50% never claim a strong match.')
p('Status: content, behavior, explicit feedback and discovery are live in the local prototype. The embedding and LLM stages are built and tested with stand-in data; they activate after build_story_embeddings.py and llm_rerank.py are run on a computer with YouTube, Hugging Face and LLM API access.')
p('Not used, with reasons: user-based collaborative filtering needs many listeners; item-based filtering from one listener’s sessions was too sparse (10 of 718 story pairs recur); two-tower, graph and generative recommenders need large-scale interaction data. Playlist-based item similarity remains a later option.')
h('D Build the first website')
table(['ID','Small step and completion evidence','Progress'],[
 ('D1','Design the homepage using Homepage Content.md: problem, vision, planned monetization and discovery actions. Sketch preferences, recommendations, details and queue. Done when the content and complete mobile journey are reviewable.','Not started'),
 ('D2','Choose the stack and implement catalog loading and browser local storage. Done when setup runs and Cloudflare Pages deployment requirements are checked.','Not started'),
 ('D3','Build preference and favorite-link input. Done when valid inputs save and invalid links receive clear feedback.','Not started'),
 ('D4','Build filtered recommendation cards with reasons. Done when catalog data and ranking rules produce usable results.','Not started'),
 ('D5','Add YouTube playback or open-on-YouTube links. Done when available and unavailable videos behave as intended.','Not started'),
 ('D6','Add feedback, saved queue and preference reset. Done when actions persist and affect later recommendations.','Not started')],[.45,5.15,1.37])
p('Dependency: catalog and baseline rules before D4; playback feasibility before D5. Login, automated listening-history import, background playback, downloads and a native app are deferred pending need and feasibility.')

doc.add_page_break()
h('E Validate and release')
table(['ID','Small step and completion evidence','Progress'],[
 ('E1','Check phone, desktop and keyboard use; verify hard filters, reviewed tags, explanations, empty matches, broken videos and preference reset. Done when critical issues are resolved.','Not started'),
 ('E2','Run a small listener trial. Proposed: five listeners, each reviewing three suggestions. Done when choices, ratings and comments are recorded.','Not started'),
 ('E3','Review results and improve weak recommendations. Done when the release criteria are met or remaining issues are explicitly accepted.','Not started'),
 ('E4','Prepare a reviewable release, privacy text and deployment instructions. Done when the website is released after approval and a smoke check passes.','Not started')],[.45,5.15,1.37])
p('Proposed trial criteria: no hard-filter violations; at least four of five listeners find one of their top three suggestions worth listening to; median time from opening recommendations to choosing a story is under two minutes. Agree or revise these targets in A4. A playback click is not evidence that a story was enjoyed.')
doc.add_page_break()
h('F Multi-user platform')
p('Small problem: how does the website learn from many listeners instead of one? Today feedback is saved only in each visitor’s browser, so the site cannot learn from its users. The personal YouTube history used for testing never ships; the public site starts from the curated catalog. Output: accounts, a shared feedback store and recommendations that improve as more people use the site.')
p('Each user keeps a personal taste profile built from their own clicks, using the same formula as Version 3. Shared parts are learned from everyone: the weights γ and δ, plot and mood similarity, popularity and, later, “listeners who liked this also liked”.')
table(['ID','Small step and completion evidence','Progress'],[
 ('F1','Build quick-start onboarding: a new visitor picks 3–5 loved stories or genres. Done when a visitor with no history gets recommendations within one minute.','Not started'),
 ('F2','Add a popularity fallback for brand-new users, smoothed so a few early likes cannot dominate. Done when an empty profile shows popular stories with a clear reason.','Not started'),
 ('F3','Add sign-in. Done when a user’s feedback follows them across devices.','Not started'),
 ('F4','Store feedback on a server: likes, not for me, heard and listens per user. Done when events are saved, readable per user and deletable on request.','Not started'),
 ('F5','Publish a privacy notice, collect consent, and support data export and deletion; review obligations under India’s DPDP Act with a qualified adviser. Done when the notice is live and deletion works end to end.','Not started'),
 ('F6','Play stories through the embedded YouTube player and record listening progress for consenting users. Done when completion is stored and used as a signal.','Not started'),
 ('F7','Refit shared weights (γ, δ and others) from all users on a schedule. Done when the leave-one-out comparison runs on multi-user data.','Not started'),
 ('F8','Add collaborative filtering (“listeners who liked this also liked”). Done when it beats the content model on held-out feedback; start after about 1,000 active users.','Not started')],[.45,5.15,1.37])
table(['Stage','Users','What to add','Platform'],[
 ('1 Collect','First 1,000','F1–F6: onboarding, popularity, sign-in, feedback database, privacy, listening time','Cloudflare Pages with Cloudflare D1, or Supabase'),
 ('2 Learn from the crowd','1,000–100,000','F7–F8: shared weights and collaborative filtering, retrained weekly','Laptop or Google Colab'),
 ('3 Scale','100,000+','Neural recommenders learning from everyone’s behavior','Google Vertex AI or AWS SageMaker')],[1.3,1.1,2.9,1.67])
p('Dependencies: F5 before any data collection goes live (F4, F6); F3 before F4; F1 and F2 are needed at public launch. Each new method is added on top of the current model and kept only if it predicts better on held-out feedback.')
h('Open decisions')
table(['Decision','Current position','Resolve by'],[
 ('Initial language and audience','To be decided','A1'),('Catalog source and YouTube integration','Research required','B2'),('AI tagging tool and preparation cost','To be decided; human review required','B5'),('Recommendation weights','Decided in Version 3: probabilistic two-stage model; γ and δ fitted from labels','C2 (Done)'),('Embedding model','multilingual-E5-base proposed; neighbor spot-check pending','Before using emb'),('LLM provider for Stage 2','Anthropic or Gemini, offline only; key kept in .env','Before llm_rerank.py'),('Hosting and storage','Proposed Cloudflare Pages and browser local storage; stack remains open','D2'),('Live AI service','Deferred beyond first version','Future review'),('Sign-in method','To be decided (for example Google sign-in or email link)','F3'),('Feedback database','Cloudflare D1 or Supabase; to be decided','F4'),('Data retention and legal review','To be decided; DPDP Act review needed','F5'),('Owners and target dates','Unassigned and unscheduled','Before each milestone')],[2.35,3.25,1.37])
h('Blockers and next actions')
table(['Issue','Next action','Status'],[
 ('Audience is not selected','Complete A1','Open'),('YouTube capabilities are unverified','Complete B2 before promising integration features','Open'),('Plot/mood embeddings not built yet','Run build_story_embeddings.py on your PC and spot-check neighbors of the favorites','Open'),('Few labels: 3 favorites, no dislikes','Label about 30 stories as loved, fine or not in research/labels.json, then rerun fit_preferences.py','Open'),('Feedback stays in each browser, so the site cannot learn across users','Complete F3–F5','Open')],[2.35,3.25,1.37])
h('Work session log')
table(['Date','Steps changed and evidence','Next action'],[
 ('4 Oct 2026','Version 1 created; website direction captured.','A1'),('4 Oct 2026','Version 2 adds AI tagging, human review and catalog export; browser ranking and proposed hosting clarified. 0 of 25 steps complete.','A1 then B5 after prerequisites'),('5 Oct 2026','Version 3 revises the recommendation logic: two-stage probabilistic model with Beta attribute beliefs, weak capped history signals, explicit-feedback overrides, plot/mood embeddings, offline LLM rerank and Thompson sampling for discovery. C2 and C4 Done; C1 and C3 In progress. Milestones A, B, D and E not re-audited.','Build embeddings, label about 30 stories, re-audit A, B, D and E'),('5 Oct 2026','Version 4 adds milestone F, the multi-user roadmap: onboarding, popularity fallback, sign-in, server-side feedback, privacy, listening time, shared weights and collaborative filtering. Checklist now 33 steps.','F5 decisions, then F1–F4'),('To fill','Record IDs, new status, evidence and blockers.','To fill')],[1.05,4.55,1.37])
doc.core_properties.title='YouTube Audio Story Website Project Tracker'
doc.core_properties.subject='Small steps and side by side progress tracking'
doc.save(ROOT/'Project Tracker.docx')
